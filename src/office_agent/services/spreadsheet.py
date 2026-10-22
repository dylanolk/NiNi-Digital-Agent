from datetime import datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path
import re
from typing import Any
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.table import Table, TableStyleInfo
from office_agent.config import settings
from office_agent.services.product_search import ProductSearchService

class SpreadsheetService:
    HEADERS = ["Category", "Manufacturer", "Product", "Model", "SKU", "Quantity", "Currency", "Unit Price", "Line Total", "Recommendation Reason", "Manual Evidence", "Manual Source", "Price As Of", "Vendor"]
    def __init__(self, products: ProductSearchService, export_root: Path | None = None):
        self.products = products; self.export_root = export_root or settings.local_storage_path / "exports"; self.export_root.mkdir(parents=True, exist_ok=True)
    def generate(self, project_name: str, client_name: str | None = None, notes: str | None = None, recommendations: list[dict[str, Any]] | None = None) -> dict[str, Any]:
        if not project_name or not project_name.strip(): raise ValueError("Project name is required")
        recommendations = recommendations or []; workbook = Workbook(); sheet = workbook.active; sheet.title = "Recommendations"; sheet.append(self.HEADERS); warnings: list[str] = []; totals: dict[str, Decimal] = {}
        for item in recommendations:
            quantity = item.get("quantity")
            if not isinstance(quantity, int) or isinstance(quantity, bool) or quantity < 1: raise ValueError("Invalid quantity")
            if not item.get("product_id"): raise ValueError("Product ID is required")
            product = self.products.get(item["product_id"]); price = product["price"]; unit_price = line_total = None; currency = None
            if price is None: warnings.append(f"Missing price: {product['name']}")
            else:
                currency = price["currency"]; unit_price = Decimal(price["amount"]); line_total = unit_price * quantity; totals[currency] = totals.get(currency, Decimal("0.00")) + line_total
                captured = datetime.fromisoformat(price["captured_at"]); captured = captured if captured.tzinfo else captured.replace(tzinfo=timezone.utc)
                if datetime.now(timezone.utc) - captured > timedelta(days=settings.price_stale_after_days): warnings.append(f"Stale price: {product['name']}")
            evidence = item.get("manual_evidence") or []; evidence_text = "; ".join(f"{entry.get('document_id')} p.{entry.get('page_number')}: {entry.get('note', '')}" for entry in evidence); source_text = "; ".join(str(entry["source_url"]) for entry in evidence if entry.get("source_url")); row = sheet.max_row + 1
            sheet.append([product["category"], product["manufacturer"], product["name"], product["model"], product["sku"], quantity, currency, unit_price, None, item.get("recommendation_reason"), evidence_text, source_text, price["captured_at"] if price else None, price["vendor"] if price else None])
            if unit_price is not None: sheet.cell(row, 9, f"=F{row}*H{row}")
        data_end = sheet.max_row; total_row = data_end + 2
        if len(totals) <= 1:
            total = next(iter(totals.values()), Decimal("0.00")); sheet.cell(total_row, 8, "TOTAL EQUIPMENT COST"); sheet.cell(total_row, 9, f"=SUM(I2:I{data_end})" if data_end > 1 else "=0")
        else:
            total = None; warnings.append("Multiple currencies present; equipment totals are not combined"); sheet.cell(total_row, 8, "TOTAL EQUIPMENT COST (SEE CURRENCY TOTALS)")
        sheet.cell(total_row, 8).font = sheet.cell(total_row, 9).font = Font(bold=True)
        if data_end > 1:
            table = Table(displayName="RecommendationsTable", ref=f"A1:N{data_end}"); table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True); sheet.add_table(table)
        sheet.freeze_panes = "A2"; sheet.auto_filter.ref = f"A1:N{data_end}"
        for cell in sheet[1]: cell.font = Font(bold=True); cell.alignment = Alignment(wrap_text=True)
        for row in sheet.iter_rows(min_row=2, max_row=data_end):
            row[7].number_format = row[8].number_format = '#,##0.00'
            for cell in row: cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column in sheet.columns:
            letter = column[0].column_letter; sheet.column_dimensions[letter].width = min(45, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        info = workbook.create_sheet("Project Info"); info.append(["Field", "Value"]); info.append(["Project name", project_name]); info.append(["Client", client_name]); info.append(["Generated timestamp", datetime.now(timezone.utc).isoformat()]); info.append(["Notes", notes]); info.append(["Warnings", "; ".join(warnings)]); info.append(["Number of items", len(recommendations)]); info.append(["Equipment total", total]); info.append(["Currency totals", "; ".join(f"{key} {value:.2f}" for key, value in totals.items())]); info[1][0].font = info[1][1].font = Font(bold=True)
        filename = re.sub(r"[^a-z0-9]+", "-", project_name.lower()).strip("-") or "recommendations"; path = self.export_root / f"{filename}-{datetime.now():%Y%m%d%H%M%S%f}.xlsx"; workbook.save(path)
        return {"filename": path.name, "total": f"{total:.2f}" if total is not None else None, "currency": next(iter(totals), None), "currency_totals": {key: f"{value:.2f}" for key, value in totals.items()}, "line_items": len(recommendations), "warnings": warnings, "download_reference": str(path)}
