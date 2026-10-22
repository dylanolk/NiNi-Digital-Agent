from datetime import datetime, timedelta, timezone
from decimal import Decimal
import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from office_agent.db.models import Base, ProductPrice
from office_agent.repositories.prices import PriceRepository
from office_agent.repositories.products import ProductRepository
from office_agent.services.catalog import CatalogService
from office_agent.services.product_search import ProductSearchService
from office_agent.services.spreadsheet import SpreadsheetService

@pytest.fixture
def spreadsheet(tmp_path):
    engine = create_engine("sqlite://"); Base.metadata.create_all(engine); session = sessionmaker(bind=engine)(); catalog = CatalogService(ProductRepository(session), PriceRepository(session)); ap = catalog.add_product("Acme", "Office AP", "AP-1", "access_point", sku="APSKU"); switch = catalog.add_product("Acme", "PoE Switch", "SW-1", "network_switch", sku="SWSKU"); catalog.add_price(ap.id, "Vendor A", "200.00"); catalog.add_price(switch.id, "Vendor A", "600.00"); return SpreadsheetService(ProductSearchService(ProductRepository(session), PriceRepository(session)), tmp_path), ap, switch

def test_workbook_contains_authoritative_lines_formulas_and_evidence(spreadsheet):
    service, ap, switch = spreadsheet; result = service.generate("Acme Office", "Acme", "Network upgrade", [{"product_id": str(ap.id), "quantity": 4, "recommendation_reason": "Coverage", "manual_evidence": [{"document_id": "doc-1", "page_number": 7, "note": "PoE+ requirement", "source_url": "https://example.test/manual.pdf"}]}, {"product_id": str(switch.id), "quantity": 1, "recommendation_reason": "Power budget"}])
    assert result["total"] == "1400.00" and not result["warnings"]
    workbook = load_workbook(result["download_reference"], data_only=False); sheet = workbook["Recommendations"]
    assert sheet["H2"].value == 200 and sheet["I2"].value == "=F2*H2" and sheet["I3"].value == "=F3*H3" and sheet["I5"].value == "=SUM(I2:I3)"
    assert sheet["J2"].value == "Coverage" and "p.7" in sheet["K2"].value and sheet["L2"].value.endswith("manual.pdf")
    assert "Project Info" in workbook.sheetnames

def test_missing_and_stale_prices_warn_and_do_not_become_zero(spreadsheet):
    service, ap, switch = spreadsheet; session = service.products.price_service.prices.session; session.query(ProductPrice).filter(ProductPrice.product_id == ap.id).update({"captured_at": datetime.now(timezone.utc) - timedelta(days=90)}); session.query(ProductPrice).filter(ProductPrice.product_id == switch.id).delete(); session.commit(); result = service.generate("Warnings", recommendations=[{"product_id": str(ap.id), "quantity": 1}, {"product_id": str(switch.id), "quantity": 2}])
    assert result["total"] == "200.00" and any("Stale price" in warning for warning in result["warnings"]) and any("Missing price" in warning for warning in result["warnings"])
    assert load_workbook(result["download_reference"], data_only=False)["Recommendations"]["H3"].value is None

def test_quantity_and_project_validation(spreadsheet):
    service, ap, _ = spreadsheet
    with pytest.raises(ValueError, match="Invalid quantity"): service.generate("Invalid", recommendations=[{"product_id": str(ap.id), "quantity": 0}])
    with pytest.raises(ValueError, match="Project name"): service.generate("", [])
